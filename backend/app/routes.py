from flask import Blueprint, request, jsonify, current_app, send_from_directory
from werkzeug.utils import secure_filename
import os
from .models import db, Job, Candidate
import re
import string

# Try to import scikit-learn, but provide fallbacks if not available
try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
    sklearn_available = True
except ImportError:
    print("scikit-learn not available. Using basic scoring methods only.")
    sklearn_available = False
    
# Update PyPDF2 import with better error handling
try:
    import PyPDF2
except ImportError:
    print("Error importing PyPDF2, trying alternative installations...")
    try:
        import pypdf as PyPDF2
    except ImportError:
        print("Failed to import any PDF library. PDF processing may fail.")
from docx import Document
from sqlalchemy.exc import SQLAlchemyError
import re
import uuid
import spacy
from nltk.corpus import stopwords
from nltk.tokenize import word_tokenize
import string

# Initialize Blueprint
main = Blueprint('main', __name__)

# Configure allowed file extensions
ALLOWED_EXTENSIONS = {'pdf', 'docx'}

# Load spaCy model
try:
    nlp = spacy.load("en_core_web_sm")
except Exception as e:
    print(f"Error loading spaCy model: {str(e)}")
    # Set nlp to None if loading fails
    nlp = None

# Initialize NLTK stopwords
try:
    import nltk
    nltk.download('stopwords', quiet=True)
    nltk.download('punkt', quiet=True)
    stop_words = set(stopwords.words('english'))
except Exception as e:
    print(f"Error initializing NLTK: {str(e)}")
    # Fallback to a simple list of common English stopwords
    stop_words = set(['i', 'me', 'my', 'myself', 'we', 'our', 'ours', 'ourselves', 'you', 'your', 'yours',
                   'yourself', 'yourselves', 'he', 'him', 'his', 'himself', 'she', 'her', 'hers',
                   'herself', 'it', 'its', 'itself', 'they', 'them', 'their', 'theirs', 'themselves',
                   'what', 'which', 'who', 'whom', 'this', 'that', 'these', 'those', 'am', 'is', 'are',
                   'was', 'were', 'be', 'been', 'being', 'have', 'has', 'had', 'having', 'do', 'does',
                   'did', 'doing', 'a', 'an', 'the', 'and', 'but', 'if', 'or', 'because', 'as', 'until',
                   'while', 'of', 'at', 'by', 'for', 'with', 'about', 'against', 'between', 'into',
                   'through', 'during', 'before', 'after', 'above', 'below', 'to', 'from', 'up', 'down',
                   'in', 'out', 'on', 'off', 'over', 'under', 'again', 'further', 'then', 'once', 'here',
                   'there', 'when', 'where', 'why', 'how', 'all', 'any', 'both', 'each', 'few', 'more',
                   'most', 'other', 'some', 'such', 'no', 'nor', 'not', 'only', 'own', 'same', 'so',
                   'than', 'too', 'very', 's', 't', 'can', 'will', 'just', 'don', 'should', 'now', 'd',
                   'll', 'm', 'o', 're', 've', 'y', 'ain', 'aren', 'couldn', 'didn', 'doesn', 'hadn',
                   'hasn', 'haven', 'isn', 'ma', 'mightn', 'mustn', 'needn', 'shan', 'shouldn', 'wasn',
                   'weren', 'won', 'wouldn'])

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_pdf(file_path):
    text = ''
    try:
        with open(file_path, 'rb') as file:
            # First try with PyPDF2 1.26.0 method
            try:
                reader = PyPDF2.PdfFileReader(file)
                # Check if PDF is encrypted
                if reader.isEncrypted:
                    try:
                        reader.decrypt('')  # Try empty password
                    except:
                        pass  # If decrypt fails, we'll still try to read what we can

                for page_num in range(reader.getNumPages()):
                    try:
                        page = reader.getPage(page_num)
                        text += page.extractText() + "\n"
                    except:
                        continue  # Skip problematic pages
            except (AttributeError, TypeError, ImportError):
                # If old method fails, try using newer PyPDF2 version methods
                try:
                    # Reset file position
                    file.seek(0)
                    # For PyPDF2 3.x
                    reader = PyPDF2.PdfReader(file)
                    for page in reader.pages:
                        try:
                            text += page.extract_text() + "\n"
                        except:
                            continue  # Skip problematic pages
                except:
                    pass  # If both methods fail, we'll return empty text
    except Exception as e:
        current_app.logger.error(f"Error reading PDF: {str(e)}")
    
    # If no text was extracted, add placeholder text
    if not text.strip():
        text = "Unable to extract text from PDF. Please check if the document contains selectable text."
    
    return text

def extract_text_from_docx(file_path):
    text = ''
    try:
        doc = Document(file_path)
        for para in doc.paragraphs:
            text += para.text + '\n'
    except Exception as e:
        current_app.logger.error(f"Error reading DOCX: {str(e)}")
    return text

def clean_text(text):
    text = re.sub(r'http\S+', '', text)
    text = re.sub(r'[^a-zA-Z0-9\s.,;:!?\-&]', '', text)
    text = text.translate(str.maketrans('', '', string.punctuation))
    text = re.sub(r'\d+', '', text)
    words = word_tokenize(text.lower())
    filtered_words = [word for word in words if word not in stop_words and len(word) > 2]
    return ' '.join(filtered_words)

def extract_keywords(text, top_n=10):
    """Extract top N keywords based on TF-IDF."""
    # If scikit-learn is not available, return simple word frequency
    if not sklearn_available:
        try:
            words = text.lower().split()
            word_freq = {}
            for word in words:
                if word not in stop_words and len(word) > 2:
                    word_freq[word] = word_freq.get(word, 0) + 1
            
            # Sort by frequency
            sorted_words = sorted(word_freq.items(), key=lambda x: x[1], reverse=True)
            return [word for word, freq in sorted_words[:top_n]]
        except Exception as e:
            current_app.logger.error(f"Error extracting keywords with basic method: {str(e)}")
            return []
    
    # If scikit-learn is available, use TF-IDF
    try:
        # Check if text is empty or too short
        if not text or len(text.split()) < 3:
            return []
            
        vectorizer = TfidfVectorizer(min_df=1, max_features=100)
        tfidf = vectorizer.fit_transform([text])
        
        # In case of empty vectorizer result
        if tfidf.shape[1] == 0:
            return []
            
        scores = zip(vectorizer.get_feature_names_out(), tfidf.toarray()[0])
        sorted_scores = sorted(scores, key=lambda x: x[1], reverse=True)
        return [word for word, score in sorted_scores[:top_n] if score > 0]
    except Exception as e:
        current_app.logger.error(f"Error extracting keywords: {str(e)}")
        return []

def calculate_score(resume_text, job_description):
    """Calculate how well a resume matches a job description."""
    try:
        # Basic text cleaning
        clean_resume = clean_text(resume_text)
        clean_job = clean_text(job_description)

        # Safety check for empty texts
        if not clean_resume or not clean_job:
            return 0.0
            
        # Basic word overlap scoring (safe method)
        resume_words = set(clean_resume.split())
        job_words = set(clean_job.split())
        
        if not job_words:
            return 0.0
            
        simple_score = len(resume_words.intersection(job_words)) / len(job_words) * 100
        
        # Only try advanced scoring if sklearn is available
        if sklearn_available and len(clean_resume.split()) > 10 and len(clean_job.split()) > 10:
            try:
                # TF-IDF with cosine similarity
                vectorizer = TfidfVectorizer(min_df=1)
                tfidf_matrix = vectorizer.fit_transform([clean_job, clean_resume])
                
                if tfidf_matrix.shape[1] > 0:  # Only if we have features
                    similarity = cosine_similarity(tfidf_matrix[0:1], tfidf_matrix[1:2])[0][0] * 100
                    
                    # Final score (70% TF-IDF similarity, 30% word overlap)
                    combined_score = (0.7 * similarity) + (0.3 * simple_score)
                    return round(min(combined_score, 100.0), 2)
            except Exception as e:
                current_app.logger.error(f"Advanced scoring failed, using simple method: {str(e)}")
                # Fall back to simple scoring if advanced fails
        
        # Return simple scoring result if advanced method was skipped or failed
        return round(min(simple_score, 100.0), 2)
    except Exception as e:
        current_app.logger.error(f"Score calculation error: {str(e)}")
        return 1.0  # Return minimal non-zero score on error

def extract_email(text):
    # More comprehensive pattern to catch various email formats
    email_patterns = [
        r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b',  # Standard email
        r'\b[A-Za-z0-9._%+-]+[\s]*@[\s]*[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b',  # Email with spaces
        r'\b[A-Za-z0-9._%+-]+[\s]*\(at\)[\s]*[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b',  # Email with (at)
        r'\b[A-Za-z0-9._%+-]+[\s]*\[at\][\s]*[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b',  # Email with [at]
    ]
    
    for pattern in email_patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            # Clean up any spaces or special formats
            email = match.group(0)
            email = email.replace(' ', '').replace('(at)', '@').replace('[at]', '@')
            return email
    
    # Try to find email mentioned in different format like "Email: user at domain.com"
    email_label_pattern = r'(?:e-?mail|id)[\s:]*([A-Za-z0-9._%+-]+[\s]*(?:@|\(at\)|\[at\])[\s]*[A-Za-z0-9.-]+\.[A-Za-z]{2,})'
    match = re.search(email_label_pattern, text, re.IGNORECASE)
    if match:
        email = match.group(1)
        email = email.replace(' ', '').replace('(at)', '@').replace('[at]', '@')
        return email
        
    return ''

def extract_phone(text):
    # Multiple patterns to catch different phone number formats
    phone_patterns = [
        r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}',  # Standard formats: (123) 456-7890, 123-456-7890
        r'\+?\d{1,3}[-.\s]?\d{2,3}[-.\s]?\d{2,4}[-.\s]?\d{2,4}',  # International format: +91 98765 43210
        r'\b\d{10}\b',  # Plain 10 digits: 1234567890
        r'\b\d{3}[-.\s]?\d{3}[-.\s]?\d{4}\b',  # Common US format: 123-456-7890
        r'\(\d{3}\)[-.\s]?\d{3}[-.\s]?\d{4}',  # Parentheses format: (123) 456-7890
        r'\b\d{5}[-.\s]?\d{5}\b',  # Format like: 12345 67890
    ]
    
    # Check each pattern
    for pattern in phone_patterns:
        matches = re.findall(pattern, text)
        if matches:
            # Clean the match by removing non-digit characters except + for international prefix
            phone = re.sub(r'[^\d+]', '', matches[0])
            # Limit to a reasonable length to avoid garbage matches
            return phone[:15]
    
    # Try to find phone mentioned after a label like "Phone: 123-456-7890"
    phone_label_pattern = r'(?:phone|mobile|cell|tel|telephone|contact)[\s:]*(\+?\d[-()\s.\d]{7,20})'
    match = re.search(phone_label_pattern, text, re.IGNORECASE)
    if match:
        phone = re.sub(r'[^\d+]', '', match.group(1))
        return phone[:15]
        
    return ''

def extract_name(text):
    # Skip NLP processing due to compatibility issues
    # Instead use a simple approach to find the name
    
    # Fallback to first line with title case
    lines = text.split('\n')
    for line in lines[:10]:  # Check just the first 10 lines
        line = line.strip()
        if line and len(line.split()) <= 4 and len(line.split()) >= 2:
            words = line.split()
            # Check if first few words look like a name (capitalized, no numbers)
            if all(word[0].isupper() for word in words if word) and not any(char.isdigit() for char in line):
                return line
    
    # Try to extract name from common formats like "Name: John Doe"
    name_patterns = [
        r'name[:\s]+([A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,3})',
        r'([A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,3})\s*(?:resume|cv)',
    ]
    
    for pattern in name_patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            return match.group(1).strip()
    
    return ''

def extract_address(text):
    # Skip NLP processing due to compatibility issues
    # Instead use regex patterns to find city names
    
    # Common city/location patterns
    city_pattern = r'(?:located|based|live|living|residing|from|in)\s+(?:in|at|near)?\s+([A-Z][a-zA-Z\s]+(?:,\s*[A-Z][a-zA-Z\s]+)*)'
    address_pattern = r'(?:address|location)(?::|is|:is)?\s+([A-Z][a-zA-Z0-9\s,.]+)(?:\.|\n|$)'
    
    # Try to find location using patterns
    for pattern in [city_pattern, address_pattern]:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            # Clean up the match
            location = match.group(1).strip()
            # Remove common non-location words
            for word in ['resume', 'cv', 'curriculum', 'vitae', 'contact', 'details', 'information']:
                location = re.sub(r'\b' + word + r'\b', '', location, flags=re.IGNORECASE)
            return location.strip()
    
    # Try to find any postal code or zip code patterns
    zipcode_patterns = [
        r'[A-Z]{1,2}[0-9][A-Z0-9]? [0-9][A-Z]{2}',  # UK postcodes
        r'\b\d{5}(?:-\d{4})?\b',  # US zipcodes
        r'[A-Z]\d[A-Z] \d[A-Z]\d'  # Canadian postcodes
    ]
    
    for pattern in zipcode_patterns:
        match = re.search(pattern, text)
        if match:
            # Try to find city name near the zipcode
            context = text[max(0, match.start() - 50):min(len(text), match.end() + 50)]
            city_near_zip = re.search(r'([A-Z][a-zA-Z\s]+)(?:,|\s+)(?:\d|\s|[A-Z])+' + re.escape(match.group(0)), context)
            if city_near_zip:
                return city_near_zip.group(1).strip()
            return match.group(0)
    
    return ''

def extract_highest_qualification(text):
    qualifications = {
        'PhD': r'\b(ph\.?d|doctorate)\b',
        'Masters': r'\b(m\.?sc|master\'s|m\.?a|m\.?ed|msc|ma|ms)\b',
        'Bachelors': r'\b(b\.?sc|bachelor\'s|b\.?a|b\.?ed|bsc|ba)\b',
        'Diploma': r'\b(diploma|certificate)\b',
        'High School': r'\b(h\.s|high school|secondary school)\b'
    }
    
    found = []
    text_lower = text.lower()
    for qual, pattern in qualifications.items():
        if re.search(pattern, text_lower):
            found.append(qual)
    
    hierarchy = ['PhD', 'Masters', 'Bachelors', 'Diploma', 'High School']
    for level in hierarchy:
        if level in found:
            return level
    return ''

@main.route('/')
def index():
    return jsonify({
        'message': 'Resume Shortlister API',
        'status': 'running',
        'endpoints': {
            'jobs': '/api/jobs',
            'upload_resume': '/api/jobs/<job_id>/upload-resume',
            'candidates': '/api/jobs/<job_id>/candidates'
        }
    })

@main.route('/api/jobs', methods=['POST'])
def create_job():
    try:
        data = request.get_json()
        if not data or 'title' not in data or 'description' not in data:
            return jsonify({'error': 'Missing required fields'}), 400

        job = Job(title=data['title'], description=data['description'])
        db.session.add(job)
        db.session.commit()
        return jsonify({'id': job.id, 'title': job.title}), 201
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500

@main.route('/api/jobs', methods=['GET'])
def get_jobs():
    try:
        jobs = Job.query.all()
        return jsonify({"jobs": [{'id': j.id, 'title': j.title, 'description': j.description} for j in jobs]})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@main.route('/api/jobs/<int:job_id>', methods=['DELETE'])
def delete_job(job_id):
    try:
        job = Job.query.get_or_404(job_id)
        
        Candidate.query.filter_by(job_id=job_id).delete()
        db.session.delete(job)
        db.session.commit()
        
        return jsonify({
            'message': 'Job deleted successfully',
            'id': job_id
        })
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500

@main.route('/api/jobs/<int:job_id>/upload-resume', methods=['POST'])
def upload_resume(job_id):
    file_path = None
    try:
        if 'file' not in request.files:
            return jsonify({'error': 'No file provided'}), 400
        
        file = request.files['file']
        if not file or file.filename == '':
            return jsonify({'error': 'Invalid file'}), 400
        
        if not allowed_file(file.filename):
            return jsonify({'error': 'Invalid file type'}), 400

        job = Job.query.get_or_404(job_id)
        original_filename = secure_filename(file.filename)
        filename = f"{uuid.uuid4().hex}_{original_filename}"
        file_path = os.path.join(current_app.config['UPLOAD_FOLDER'], filename)
        file.save(file_path)

        # Try to extract text, but don't let extraction failures stop the process
        resume_text = ""
        try:
            if filename.lower().endswith('.pdf'):
                resume_text = extract_text_from_pdf(file_path)
            else:
                resume_text = extract_text_from_docx(file_path)

            # Log the first 100 characters of extracted text for debugging
            current_app.logger.info("Extracted text preview: {}".format(resume_text[:100].replace('\n', ' ') + "..."))
        except Exception as e:
            current_app.logger.error(f"Text extraction error but continuing: {str(e)}")
            resume_text = "Error extracting text from document. Processing with minimal information."

        # Process information - with fallbacks for each extraction
        try:
            score = calculate_score(resume_text, job.description)
        except Exception as e:
            current_app.logger.error(f"Score calculation error but continuing: {str(e)}")
            score = 1.0  # Default minimal score
            
        try:
            name = extract_name(resume_text) or request.form.get('name', '')
        except Exception:
            name = request.form.get('name', '')
            
        try:
            email = extract_email(resume_text) or request.form.get('email', '')
        except Exception:
            email = request.form.get('email', '')
            
        try:
            mobile = extract_phone(resume_text) or request.form.get('mobile', '')
        except Exception:
            mobile = request.form.get('mobile', '')
            
        try:
            city = extract_address(resume_text) or request.form.get('city', '')
        except Exception:
            city = request.form.get('city', '')
            
        try:
            qualification = extract_highest_qualification(resume_text)
        except Exception:
            qualification = ""

        # Log extracted information for debugging
        current_app.logger.info(f"Extracted info - Name: '{name}', Email: '{email}', Mobile: '{mobile}', City: '{city}'")

        # Validate required fields - only required if both extracted and form values are empty
        form_email = request.form.get('email', '')
        form_mobile = request.form.get('mobile', '')
        form_name = request.form.get('name', '')
        
        # Use provided values if extraction failed
        if not email and form_email:
            email = form_email
        if not mobile and form_mobile:
            mobile = form_mobile
        if not name and form_name:
            name = form_name
        
        # Generate a name from the filename if no name was found
        if not name:
            clean_filename = original_filename.replace('.pdf', '').replace('.docx', '')
            name = clean_filename.replace('_', ' ').title()
        
        # Generate placeholder email if needed
        if not email and not mobile:
            import hashlib
            name_hash = hashlib.md5(name.encode()).hexdigest()[:8]
            email = f"candidate_{name_hash}@placeholder.com"
            
        # ALWAYS ACCEPT THE RESUME - we've generated placeholders as needed

        # Check for existing candidate - but don't let this stop us
        try:
            existing = None
            if email:
                existing = Candidate.query.filter(
                    Candidate.email == email,
                    Candidate.job_id == job_id
                ).first()
            
            if not existing and mobile:
                existing = Candidate.query.filter(
                    Candidate.mobile == mobile,
                    Candidate.job_id == job_id
                ).first()
                
            # Log whether we found an existing candidate
            if existing:
                current_app.logger.info(f"Found existing candidate with ID {existing.id}, updating")
                
                # Update the existing candidate
                existing.name = name
                existing.mobile = mobile or existing.mobile
                existing.email = email or existing.email
                existing.city = city
                existing.highest_qualification = qualification
                existing.resume_path = filename
                existing.score = score
                db.session.commit()
                
                return jsonify({
                    'message': 'Candidate updated',
                    'candidate_id': existing.candidate_id,
                    'score': score,
                    'qualification': qualification,
                    'extracted_info': {
                        'name': name,
                        'email': email,
                        'mobile': mobile,
                        'city': city,
                        'qualification': qualification
                    }
                }), 200
        except Exception as e:
            current_app.logger.error(f"Error checking for existing candidate but continuing: {str(e)}")
            # Continue with creating a new candidate

        # Create new candidate - simplified to reduce failure points
        try:
            candidate = Candidate(
                name=name,
                email=email,
                mobile=mobile,
                city=city,
                highest_qualification=qualification,
                resume_path=filename,
                score=score,
                job_id=job_id
            )
            db.session.add(candidate)
            db.session.commit()
            
            current_app.logger.info(f"Successfully created new candidate with ID {candidate.id}")

            return jsonify({
                'message': 'Resume uploaded successfully',
                'candidate_id': candidate.candidate_id,
                'score': score,
                'qualification': qualification,
                'extracted_info': {
                    'name': name,
                    'email': email,
                    'mobile': mobile,
                    'city': city,
                    'qualification': qualification
                }
            }), 201
        except Exception as e:
            db.session.rollback()
            current_app.logger.error(f"Error creating candidate: {str(e)}")
            # This is a serious error we can't recover from
            raise

    except Exception as e:
        db.session.rollback()
        if file_path and os.path.exists(file_path):
            try:
                os.remove(file_path)
            except:
                pass
        # Log detailed error information
        import traceback
        error_details = traceback.format_exc()
        current_app.logger.error(f"Resume upload error: {str(e)}\n{error_details}")
        
        # Return a more helpful error message
        return jsonify({
            'error': f"Failed to process resume: {str(e)}",
            'details': "The server encountered an error while processing your resume. Please try again or use a different file format."
        }), 500

@main.route('/api/jobs/<int:job_id>/candidates', methods=['GET'])
def get_candidates(job_id):
    try:
        candidates = Candidate.query.filter_by(job_id=job_id).order_by(Candidate.score.desc()).all()
        return jsonify({"candidates": [{
            'id': c.id,
            'candidate_id': c.candidate_id,
            'name': c.name,
            'email': c.email,
            'mobile': c.mobile,
            'city': c.city,
            'qualification': c.highest_qualification,
            'score': c.score,
            'resume_path': c.resume_path
        } for c in candidates]})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@main.route('/api/top-resumes', methods=['GET'])
def get_top_resumes():
    try:
        limit = request.args.get('limit', default=10, type=int)
        job_id = request.args.get('job_id', default=None, type=int)
        
        query = db.session.query(
            Candidate, 
            Job.title.label('job_title')
        ).join(
            Job, 
            Candidate.job_id == Job.id
        )
        
        # Filter by job_id if provided
        if job_id:
            query = query.filter(Candidate.job_id == job_id)
            
        # Apply limit and ordering
        query = query.order_by(
            Candidate.score.desc()
        ).limit(limit)
        
        results = query.all()
        
        return jsonify({
            'top_resumes': [{
                'id': c.Candidate.id,
                'candidate_id': c.Candidate.candidate_id,
                'name': c.Candidate.name,
                'email': c.Candidate.email,
                'mobile': c.Candidate.mobile,
                'city': c.Candidate.city,
                'qualification': c.Candidate.highest_qualification,
                'score': c.Candidate.score,
                'resume_path': c.Candidate.resume_path,
                'job_id': c.Candidate.job_id,
                'job_title': c.job_title
            } for c in results]
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@main.route('/api/resumes/<filename>', methods=['GET'])
def download_resume(filename):
    try:
        if '..' in filename or filename.startswith('/'):
            return jsonify({'error': 'Invalid filename'}), 400
            
        file_path = os.path.join(current_app.config['UPLOAD_FOLDER'], filename)
        
        if not os.path.exists(file_path):
            return jsonify({'error': 'File not found'}), 404
            
        content_type = 'application/pdf' if filename.lower().endswith('.pdf') else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        
        disposition_type = request.args.get('download', 'false').lower() == 'true'
        
        return send_from_directory(
            current_app.config['UPLOAD_FOLDER'],
            filename,
            mimetype=content_type,
            as_attachment=disposition_type,
            download_name=filename.split('_', 1)[1] if '_' in filename else filename
        )
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# Error handler
@main.app_errorhandler(404)
def not_found(error):
    return jsonify({'error': 'Resource not found'}), 404